"""Blind IRAC scoring of the 4 methods on the 40 selected questions (rubric: docs/irac_rubric.md).

    PYTHONPATH=src .venv_diag/Scripts/python -m diagnosis.irac_judge            # LLM judge + human sheet
    PYTHONPATH=src .venv_diag/Scripts/python -m diagnosis.irac_judge agreement  # after the human sheet is filled

Answers are anonymised (A1..A160, method names removed) and shuffled with a fixed seed; the judge
(gpt-4.1, not the GPT-4o baseline) scores one answer per call against the gold conclusion and legal basis.
Writes experiments/e2e/irac_scores.json and experiments/e2e/irac_human_sheet.xlsx (20 questions = 80
answers, same anonymous ids; the key mapping ids to methods stays in irac_key.json).
"""

from __future__ import annotations

import json
import os
import random
import sys
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "experiments" / "e2e"
JUDGE = "gpt-4.1"
RUBRIC = "v1"  # v1 = pre-registered; v2 = R made precise after the first scoring (reported separately, all methods)
R_LINE = {
    "v1": "- R (Rule): đối chiếu với CĂN CỨ PHÁP LÝ CHUẨN (chấp nhận điều khoản tương đương nếu hợp lệ và còn hiệu lực). 1 = viện dẫn đúng, đầy đủ; 0.5 = đúng luật nhưng thiếu điều khoản then chốt hoặc thiếu văn bản hướng dẫn, hoặc lẫn điều khoản không áp dụng; 0 = sai điều luật, luật đã hết hiệu lực/bị thay thế, hoặc thiếu các điều khoản then chốt.",
    "v2": "- R (Rule): đối chiếu với CĂN CỨ PHÁP LÝ CHUẨN; chấp nhận điều khoản tương đương chỉ khi thuộc văn bản đang có hiệu lực. 1 = nêu đúng TÊN/SỐ văn bản và đúng SỐ ĐIỀU (khoản) then chốt như căn cứ chuẩn, đầy đủ; 0.5 = nêu đúng văn bản và ít nhất một điều then chốt đúng nhưng còn thiếu điều khoản then chốt khác hoặc văn bản hướng dẫn; 0 = không nêu số điều cụ thể nào (chỉ nói chung chung như \"theo Luật ...\"), hoặc dẫn sai điều, hoặc dẫn văn bản đã hết hiệu lực/bị thay thế.",
}
SEED = 20261020
W = {"I": 0.10, "R": 0.35, "A": 0.40, "C": 0.15}
GOLD_TXT = {"true": "CÓ", "false": "KHÔNG", "insufficient": "CHƯA ĐỦ DỮ KIỆN"}

JUDGE_PROMPT = """Bạn là giám khảo chấm câu trả lời tư vấn pháp luật Việt Nam theo rubric IRAC. Chấm khách quan, chỉ dựa vào rubric và đáp án chuẩn.

RUBRIC (mỗi thành phần chỉ được 0, 0.5 hoặc 1):
- I (Issue): 1 = nêu chính xác vấn đề pháp lý cần giải quyết; 0.5 = xác định vấn đề nhưng còn chung chung, chưa vào trọng tâm; 0 = lạc đề hoặc không tìm ra vấn đề cốt lõi.
{rule_line}
- A (Application): 1 = đối chiếu từng dữ kiện với điều kiện của quy định, xử lý ngoại lệ, giải thích vì sao áp dụng/không áp dụng; khi thiếu dữ kiện thì chỉ ra điều kiện nào chưa xác định được; 0.5 = có đối chiếu nhưng chưa đầy đủ hoặc chưa thuyết phục (bỏ sót điều kiện/ngoại lệ); 0 = chỉ lặp lại tình huống và điều luật, hoặc lập luận sai.
- C (Conclusion): 1 = kết luận trùng ĐÁP ÁN CHUẨN (CÓ / KHÔNG / CHƯA ĐỦ DỮ KIỆN; nếu chưa đủ dữ kiện thì nêu cần bổ sung gì) và nhất quán với phần A; 0.5 = đúng một phần hoặc đúng nhưng mơ hồ; 0 = sai hoặc không trả lời trực tiếp.
Không chấm văn phong. Độ dài câu trả lời không phải tiêu chí.

CÂU HỎI CỦA NGƯỜI DÙNG:
{question}

ĐÁP ÁN CHUẨN: {gold}{missing}
CĂN CỨ PHÁP LÝ CHUẨN:
{basis}

CÂU TRẢ LỜI CẦN CHẤM:
[I] {issue}
[R] {rule}
[A] {application}
[C] {conclusion}

Chỉ trả về JSON: {{"I": 0|0.5|1, "R": 0|0.5|1, "A": 0|0.5|1, "C": 0|0.5|1, "reason": "<một câu ngắn>"}}"""


def basis_text(b: list[dict]) -> str:
    return "\n".join(f"- {x.get('doc', '')}, Điều {x.get('article', '')}" + (f" khoản {x['clause']}" if x.get("clause") else "")
                     + (f" điểm {x['point']}" if x.get("point") else "") for x in b) or "- (không có)"


def snap(v) -> float:
    try:
        v = float(v)
    except (TypeError, ValueError):
        return 0.0
    return min((0.0, 0.5, 1.0), key=lambda x: abs(x - v))


SET = "40"  # "40" = stratified sample (irac_answers.json); "all" = every question (irac_answers_all.json)


def build_entries() -> tuple[list[dict], dict]:
    data = json.loads((OUT / ("irac_answers.json" if SET == "40" else "irac_answers_all.json")).read_text(encoding="utf-8"))
    entries = []
    for q in data:
        for method, ans in q["answers"].items():
            entries.append({"question_id": q["id"], "method": method, "q": q, "ans": ans})
    rng = random.Random(SEED)
    rng.shuffle(entries)
    for n, e in enumerate(entries, 1):
        e["anon"] = f"A{n:03d}" if SET == "40" else f"B{n:04d}"
    key = {e["anon"]: {"question_id": e["question_id"], "method": e["method"]} for e in entries}
    return entries, key


def cache_key(e: dict) -> str:
    """Score cache key: the judged content itself (question, gold, answer text), the judge and the rubric — so a
    changed answer is always re-scored, while an identical answer reuses its score."""
    import hashlib

    blob = json.dumps([e["q"]["question"], e["q"]["gold_expected"], e["ans"]], ensure_ascii=False, sort_keys=True)
    return f"{hashlib.sha1(blob.encode('utf-8')).hexdigest()}|{JUDGE}|{RUBRIC}"


def judge_all(entries: list[dict]) -> dict:
    load_dotenv(ROOT / ".env")
    client = OpenAI(api_key=os.environ["LLM_API_KEY"], base_url=os.environ.get("LLM_BASE_URL") or None)
    cache_f = OUT / "irac_judge_cache.json"
    cache = json.loads(cache_f.read_text(encoding="utf-8")) if cache_f.exists() else {}

    def prompt(e):
        q, a = e["q"], e["ans"]
        miss = f" (thông tin cần bổ sung: {', '.join(q['gold_missing'])})" if q["gold_expected"] == "insufficient" and q["gold_missing"] else ""
        return JUDGE_PROMPT.format(rule_line=R_LINE[RUBRIC], question=q["question"], gold=GOLD_TXT[q["gold_expected"]], missing=miss,
                                   basis=basis_text(q["gold_legal_basis"]), issue=a.get("issue", ""), rule=a.get("rule", ""),
                                   application=a.get("application", ""), conclusion=a.get("conclusion", ""))

    def call(e):
        k = cache_key(e)
        if k in cache:
            return k, cache[k]
        for attempt in range(8):
            try:
                r = client.chat.completions.create(model=JUDGE, temperature=0, seed=0, response_format={"type": "json_object"},
                                                   messages=[{"role": "user", "content": prompt(e)}])
                return k, json.loads(r.choices[0].message.content)
            except Exception as ex:
                if "insufficient_quota" in str(ex):  # out of credits: stop, keep what is done
                    raise RuntimeError("OpenAI credits exhausted; top up and re-run (scored answers are cached)") from ex
                time.sleep(min(60, 3 * 2 ** attempt))
        raise RuntimeError("judge failed")

    try:
        with ThreadPoolExecutor(12) as pool:
            for n, (k, v) in enumerate(pool.map(call, entries), 1):
                cache[k] = v
                if n % 20 == 0:
                    cache_f.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    finally:
        cache_f.write_text(json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8")
    return {e["anon"]: cache[cache_key(e)] for e in entries}


def human_sheet(entries: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    qids = sorted({e["question_id"] for e in entries})
    rng = random.Random(SEED + 1)
    pick = set(rng.sample(qids, 20))
    rows = sorted([e for e in entries if e["question_id"] in pick], key=lambda e: e["anon"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Chấm IRAC"
    hdr = ["Mã câu trả lời", "Câu hỏi", "Đáp án chuẩn", "Căn cứ pháp lý chuẩn", "[I] Vấn đề", "[R] Quy định", "[A] Áp dụng",
           "[C] Kết luận", "I (0/0.5/1)", "R (0/0.5/1)", "A (0/0.5/1)", "C (0/0.5/1)", "Ghi chú"]
    ws.append(hdr)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for e in rows:
        q, a = e["q"], e["ans"]
        miss = f" — cần bổ sung: {', '.join(q['gold_missing'])}" if q["gold_expected"] == "insufficient" and q["gold_missing"] else ""
        ws.append([e["anon"], q["question"], GOLD_TXT[q["gold_expected"]] + miss, basis_text(q["gold_legal_basis"]),
                   a.get("issue", ""), a.get("rule", ""), a.get("application", ""), a.get("conclusion", ""), None, None, None, None, ""])
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLM", [10, 50, 18, 35, 30, 40, 55, 35, 8, 8, 8, 8, 25]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    g = wb.create_sheet("Rubric")
    for line in (ROOT / "docs" / "irac_rubric.md").read_text(encoding="utf-8").splitlines():
        g.append([line])
    g.column_dimensions["A"].width = 160
    wb.save(OUT / "irac_human_sheet.xlsx")


def summarise(scores: dict, key: dict) -> dict:
    by = defaultdict(list)
    for anon, s in scores.items():
        comp = {c: snap(s.get(c)) for c in W}
        by[key[anon]["method"]].append(comp)
    out = {}
    for m, lst in by.items():
        n = len(lst)
        comp = {c: round(sum(x[c] for x in lst) / n, 3) for c in W}
        out[m] = {"n": n, **comp, "IRAC_0_100": round(100 * sum(W[c] * comp[c] for c in W), 1)}
    return out


def agreement(key: dict, scores: dict) -> None:
    from openpyxl import load_workbook

    ws = load_workbook(OUT / "irac_human_sheet.xlsx")["Chấm IRAC"]
    pairs = defaultdict(list)
    tot_h, tot_j = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        anon, hv = r[0], r[8:12]
        if anon is None or any(v is None for v in hv):
            continue
        j = {c: snap(scores[anon].get(c)) for c in W}
        h = dict(zip(W, (snap(v) for v in hv)))
        for c in W:
            pairs[c].append((h[c], j[c]))
        tot_h.append(sum(W[c] * h[c] for c in W))
        tot_j.append(sum(W[c] * j[c] for c in W))

    def kappa(ps):
        labs = (0.0, 0.5, 1.0)
        n = len(ps)
        po = sum(a == b for a, b in ps) / n
        pe = sum((sum(a == l for a, _ in ps) / n) * (sum(b == l for _, b in ps) / n) for l in labs)
        return (po - pe) / (1 - pe) if pe < 1 else 1.0

    import statistics

    res = {c: {"n": len(ps), "exact_agreement": round(sum(a == b for a, b in ps) / len(ps), 3), "kappa": round(kappa(ps), 3)}
           for c, ps in pairs.items() if ps}
    if len(tot_h) > 2:
        res["total_pearson"] = round(statistics.correlation(tot_h, tot_j), 3)
    (OUT / "irac_agreement.json").write_text(json.dumps(res, indent=1), encoding="utf-8")
    print(json.dumps(res, indent=1))


def main() -> None:
    global RUBRIC, SET
    sys.stdout.reconfigure(encoding="utf-8")
    if "--all" in sys.argv:
        SET = "all"
    if "--rubric" in sys.argv:
        RUBRIC = sys.argv[sys.argv.index("--rubric") + 1]
    suffix = ("" if RUBRIC == "v1" else f"_{RUBRIC}") + ("" if SET == "40" else "_all")
    entries, key = build_entries()
    if len(sys.argv) > 1 and sys.argv[1] == "agreement":
        scores = json.loads((OUT / f"irac_scores{suffix}.json").read_text(encoding="utf-8"))["scores"]
        agreement(key, scores)
        return
    (OUT / ("irac_key.json" if SET == "40" else "irac_key_all.json")).write_text(json.dumps(key, indent=1), encoding="utf-8")
    scores = judge_all(entries)
    summ = summarise(scores, key)
    (OUT / f"irac_scores{suffix}.json").write_text(json.dumps({"judge": JUDGE, "rubric": RUBRIC, "summary": summ, "scores": scores},
                                                              ensure_ascii=False, indent=1), encoding="utf-8")
    if RUBRIC == "v1" and SET == "40":
        human_sheet(entries)
    print(json.dumps(summ, indent=1))


if __name__ == "__main__":
    main()
