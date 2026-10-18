"""Merge LLM (Claude) review labels, build the reviewed KB, export a human-review workbook, score agreement.

    PYTHONPATH=src .venv_diag/Scripts/python -m diagnosis.kbx.review merge
    PYTHONPATH=src .venv_diag/Scripts/python -m diagnosis.kbx.review agreement   (after the human fills the workbook)
"""

from __future__ import annotations

import json
import random
import re
import sys
from collections import Counter
from pathlib import Path

from .corpus import ROOT

DIR = ROOT / "data" / "diag" / "kb_full_draft"
REV = DIR / "review"
XLSX = REV / "human_review_sample.xlsx"
HUMAN_COLS = ["Nhãn của bạn (correct/fix/drop)", "Rule đúng (nếu khác)", "Ghi chú"]


def _load_batches() -> dict[str, dict]:
    rules = {}
    for f in sorted((REV / "batches").glob("batch_*.md")):
        for block in f.read_text(encoding="utf-8").split("######## ")[1:]:
            head, _, rest = block.partition("\n")
            text, _, rule_lines = rest.rpartition("[RULES]")
            text = text.replace("[NGUYÊN VĂN]", "").strip()
            for line in rule_lines.splitlines():
                if line.startswith("{"):
                    r = json.loads(line)
                    r.update(batch=f.stem, article_head=head.strip(), article_text=text)
                    rules[r["id"]] = r
    return rules


def _load_labels() -> dict[str, dict]:
    labels = {}
    for f in sorted((REV / "out").glob("batch_*.jsonl")):
        for line in f.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line:
                d = json.loads(line)
                labels[d["id"]] = d
    return labels


def fmt(premises, conclusion) -> str:
    return f"{' ∧ '.join(premises) if premises else '(không điều kiện)'} ⇒ {conclusion}"


def merge() -> None:
    rules, labels = _load_batches(), _load_labels()
    missing = [i for i in rules if i not in labels]
    extra = [i for i in labels if i not in rules]
    verdicts = Counter(labels[i]["verdict"] for i in rules if i in labels)
    conf = Counter(labels[i].get("confidence", "high") for i in rules if i in labels)

    reviewed = []
    for i, r in rules.items():
        lab = labels.get(i)
        if lab is None or lab["verdict"] == "drop":
            continue
        out = dict(r) if lab["verdict"] == "correct" else dict(
            r, premises=lab.get("premises", r["premises"]), conclusion=lab.get("conclusion", r["conclusion"]),
            source=lab.get("source", r["source"]), exception_to=lab.get("exception_to", r.get("exception_to", "")))
        for k in ("batch", "article_head", "article_text"):
            out.pop(k, None)
        out.update(review_status=f"llm_reviewed:{lab['verdict']}", review_confidence=lab.get("confidence", "high"),
                   review_reason=lab.get("reason", ""))
        reviewed.append(out)
    with open(DIR / "rules_bhxh_reviewed.jsonl", "w", encoding="utf-8") as fh:
        for r in reviewed:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")

    # human workbook: every low-confidence rule + a stratified random sample of the rest
    rng = random.Random(42)
    low = [i for i in rules if labels.get(i, {}).get("confidence") == "low"]
    rest = [i for i in rules if i in labels and i not in low]
    sample = []
    for v in ("correct", "fix", "drop"):
        pool = [i for i in rest if labels[i]["verdict"] == v]
        k = {"correct": 70, "fix": 50, "drop": 30}[v]
        sample += rng.sample(pool, min(k, len(pool)))
    chosen = [("low_confidence", i) for i in low] + [("random_sample", i) for i in sample]

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    head = ["#", "Lý do chọn", "Rule id", "Văn bản / Điều", "Nguyên văn điều luật (rút gọn)", "Rule gốc (gpt-4o-mini)",
            "Source gốc", "Nhãn của Claude", "Rule sau khi Claude sửa", "Source sau sửa", "Lý do của Claude",
            "Độ tự tin", *HUMAN_COLS]
    ws.append(head)
    for n, (why, i) in enumerate(chosen, 1):
        r, lab = rules[i], labels[i]
        text = r.get("article_text", "")
        ws.append([n, why, i, r["article_head"][:120], text[:3000], fmt(r["premises"], r["conclusion"]),
                   r["source"].split(", ", 1)[-1], lab["verdict"],
                   fmt(lab.get("premises", []), lab.get("conclusion", "")) if lab["verdict"] == "fix" else "",
                   lab.get("source", "") if lab["verdict"] == "fix" else "", lab.get("reason", ""),
                   lab.get("confidence", "high"), "", "", ""])
    widths = [5, 14, 26, 30, 70, 45, 22, 10, 45, 22, 40, 9, 16, 40, 30]
    for col, w in zip("ABCDEFGHIJKLMNO", widths):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[12:15]:
            c.fill = fill
    ws.freeze_panes = "C2"
    guide = wb.create_sheet("hướng dẫn")
    for line in [
        "Điền 3 cột màu vàng (M, N, O) cho mỗi dòng.",
        "Cột M: nhãn ĐÚNG theo bạn cho RULE GỐC: correct (rule gốc đúng) / fix (rule gốc sai nhưng sửa được) / drop (bỏ).",
        "Cột N: nếu nhãn fix, ghi rule đúng (hoặc ghi 'như Claude' nếu bản sửa của Claude đã đúng).",
        "Cột O: ghi chú tự do.",
        "Dòng 'low_confidence' là những rule Claude không chắc; dòng 'random_sample' dùng để đo độ chính xác của Claude.",
        "Xong thì lưu file và báo để chạy: python -m diagnosis.kbx.review agreement",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    wb.save(XLSX)

    stats = {
        "scope": "Luật BHXH 41/2024 + NĐ 158/2025",
        "rules_in_scope": len(rules), "labelled": len(rules) - len(missing), "missing_labels": len(missing),
        "unknown_ids_in_output": len(extra), "verdicts": dict(verdicts), "confidence": dict(conf),
        "rules_after_review": len(reviewed), "human_workbook_rows": len(chosen),
        "human_workbook_low_confidence": len(low), "human_workbook_random_sample": len(sample),
        "reviewer": "Claude (LLM) — not a human review; human agreement measured on the workbook sample",
    }
    (REV / "review_stats.json").write_text(json.dumps(stats, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(stats, ensure_ascii=False, indent=1))
    if missing:
        print("missing labels (first 10):", missing[:10])


def agreement() -> None:
    from openpyxl import load_workbook

    ws = load_workbook(XLSX).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    pairs = [(str(r[7]).strip().lower(), str(r[12]).strip().lower(), r[1]) for r in rows
             if r[12] and str(r[12]).strip().lower() in ("correct", "fix", "drop")]
    if not pairs:
        sys.exit("No human labels found in column M yet.")
    labs = ["correct", "fix", "drop"]
    n = len(pairs)
    po = sum(a == b for a, b, _ in pairs) / n
    pe = sum((sum(a == l for a, _, _ in pairs) / n) * (sum(b == l for _, b, _ in pairs) / n) for l in labs)
    kappa = (po - pe) / (1 - pe) if pe < 1 else 1.0
    rnd = [(a, b) for a, b, w in pairs if w == "random_sample"]
    out = {
        "human_labelled": n, "agreement": po, "cohen_kappa": kappa,
        "agreement_random_sample": sum(a == b for a, b in rnd) / max(1, len(rnd)),
        # precision of the draft extractor on the random sample, as judged by the human
        "draft_rule_correct_rate_random_sample": sum(b == "correct" for _, b in rnd) / max(1, len(rnd)),
        "confusion_claude_vs_human": {a: dict(Counter(b for x, b, _ in pairs if x == a)) for a in labs},
    }
    (REV / "agreement.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    {"merge": merge, "agreement": agreement}[sys.argv[1] if len(sys.argv) > 1 else "merge"]()
